#!/usr/bin/env python3
"""Recalculate an .xlsx with headless LibreOffice and report formula errors as JSON.

Usage: python recalc.py model.xlsx [timeout_seconds]

LibreOffice computes every formula while converting a copy; the copy (now
holding cached values, so validators and viewers see numbers) replaces the
original, and is scanned for Excel error values. Formulas are kept.
Exit code 0 = success, 1 = errors found, 2 = could not run.
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


def find_soffice():
    for name in ("soffice", "libreoffice"):
        path = shutil.which(name)
        if path:
            return path
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def scan(path):
    import openpyxl

    formulas = openpyxl.load_workbook(path, data_only=False)
    total = sum(
        1 for ws in formulas for row in ws.iter_rows() for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    )
    return total


def errors_in(path):
    import openpyxl

    values = openpyxl.load_workbook(path, data_only=True)
    found = {}
    for ws in values:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in ERRORS:
                    found.setdefault(c.value.strip(), []).append(f"{ws.title}!{c.coordinate}")
    return found


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    src = Path(sys.argv[1]).resolve()
    timeout = int(sys.argv[2]) if len(sys.argv) > 2 else 60
    if not src.exists():
        print(json.dumps({"status": "error", "message": f"not found: {src}"}))
        return 2
    soffice = find_soffice()
    if not soffice:
        print(json.dumps({"status": "error", "message": "LibreOffice not installed (brew install --cask libreoffice / apt install libreoffice-calc)"}))
        return 2
    with tempfile.TemporaryDirectory() as tmp:
        cmd = [soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, str(src)]
        try:
            subprocess.run(cmd, check=True, capture_output=True, timeout=timeout)
        except (subprocess.SubprocessError, OSError) as exc:
            print(json.dumps({"status": "error", "message": f"LibreOffice failed: {exc}"}))
            return 2
        out = Path(tmp) / (src.stem + ".xlsx")
        found = errors_in(out)
        total = scan(out)
        shutil.copyfile(out, src)
    total_errors = sum(len(v) for v in found.values())
    result = {"status": "errors_found" if total_errors else "success",
              "total_errors": total_errors, "total_formulas": total}
    if found:
        result["error_summary"] = {k: {"count": len(v), "locations": v[:50]} for k, v in found.items()}
    print(json.dumps(result, indent=2))
    return 1 if total_errors else 0


if __name__ == "__main__":
    sys.exit(main())
